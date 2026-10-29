# -*- coding: utf-8 -*-
"""
Created on Thu Mar 31 13:18:11 2022

@author: changkk
"""

import pandas as pd
import numpy as np
import warnings
import os
import csv
import sys
from pathlib import Path
from PyPDF2 import PdfReader
import shutil
from tempfile import TemporaryDirectory
from urllib.parse import quote, urlparse

from io import StringIO
from pdfminer.high_level import extract_text
from pdfminer.high_level import extract_pages
from pdfminer.layout import LTTextContainer, LTTextBox
from pdfminer.converter import TextConverter
from pdfminer.layout import LAParams
from pdfminer.pdfdocument import PDFDocument
from pdfminer.pdfinterp import PDFResourceManager, PDFPageInterpreter
from pdfminer.pdfpage import PDFPage
from pdfminer.pdfparser import PDFParser
from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[2]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from shared.logging import setup_logging
from shared.sharepoint.graph import GraphClient

logger = setup_logging("ro_pdf_scraper")

warnings.filterwarnings('ignore')
pd.set_option('display.width', 500)
pd.set_option('display.max_column', 500)
pd.set_option('display.max_rows', 200)
pd.set_option('display.max_colwidth', 500)

#%%file manager functions
def read_txt(filename):  # Read a .txt file
    txt_data = []
    file = open(filename)
    while True:
        text = file.readline().strip('\n')
        if text == '':
            break
        txt_data.append(text)
    return file, txt_data


def read_csv(filename, dictread=False):  # Read a .csv file, with or without DictReader
    file = open(filename, 'r', newline='', encoding='utf-8')
    if dictread:
        csv_data = csv.DictReader(file, dialect='excel')
    else:
        csv_data = csv.reader(file, dialect='excel')
    return file, csv_data

# def write_txt(filename, permission): # Write to a .txt file
#     # filename =
#     file = open(filename, permission, newline='', encoding='utf-8')


def write_csv(filename, permission, folderpath=''):  # Write to a .csv file
    filename = filename.replace("/", "-")  # Remove "/" to ensure file can be saved
    file = open(folderpath + filename, permission, newline='', encoding='utf-8')
    csv_writer = csv.writer(file, dialect='excel')
    return file, csv_writer


def write_csv_dict(filename, permission, headers, folderpath=''): # Write to a .csv file with DictWriter
    filename = filename.replace("/", "-")  # Remove "/" to ensure file can be saved
    file = open(folderpath + filename, permission, newline='', encoding='utf-8')
    csv_writer = csv.DictWriter(file, fieldnames=headers)
    return file, csv_writer

def text_parser(search_text, start_word, last_word):  # Get text in between given strings
    start_index = search_text.lower().index(start_word)
    end_index = search_text.lower().index(last_word)
    parsed_text = search_text[(start_index+len(start_word)):end_index]
    return parsed_text


def resolve_sharepoint_folder(graph, site_url, library_name, folder_path):
    parsed_url = urlparse(site_url)
    site = graph.get(f"sites/{parsed_url.netloc}:{quote(parsed_url.path, safe='/')}")
    drives = graph.get_all(f"sites/{site['id']}/drives")
    drive = next(
        (item for item in drives if item['name'].casefold() == library_name.casefold()),
        None,
    )
    if drive is None:
        raise ValueError(
            f"SharePoint library '{library_name}' not found. "
            f"Available libraries: {[item['name'] for item in drives]}"
        )
    drive_endpoint = f"drives/{drive['id']}"
    folder_path = folder_path.strip('/')
    folder_endpoint = f"{drive_endpoint}/root"
    if folder_path:
        folder_endpoint += f":/{quote(folder_path, safe='/')}"
    folder = graph.get(folder_endpoint)
    folder_endpoint = f"{drive_endpoint}/items/{folder['id']}"
    new_folder = graph.get(f"{folder_endpoint}:/New")
    old_folder = graph.get(f"{folder_endpoint}:/Old")
    return drive_endpoint, folder_endpoint, new_folder['id'], old_folder['id']


def download_sharepoint_file(graph, item_endpoint, local_path):
    item = graph.get(item_endpoint)
    download_url = item['@microsoft.graph.downloadUrl']
    with graph.session.get(download_url, stream=True, timeout=graph.timeout) as response:
        response.raise_for_status()
        with open(local_path, 'wb') as output:
            for chunk in response.iter_content(chunk_size=1024 * 1024):
                output.write(chunk)


def upload_sharepoint_workbook(graph, item_endpoint, local_path):
    graph.ensure_authenticated()
    headers = dict(graph.headers)
    headers['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    with open(local_path, 'rb') as workbook:
        response = graph.session.put(
            f"{graph.GRAPH_BASE_URL}/{item_endpoint}/content",
            headers=headers,
            data=workbook,
            timeout=graph.timeout,
        )
    response.raise_for_status()

        # Open and read config files
#Test in Laptop
wd_keywords_path = PROJECT_ROOT / 'config' / 'wd_keywords.txt'
wd_keywords_file = read_txt(wd_keywords_path)
wd_keywords = wd_keywords_file[1]  # Keywords for identifying potential wire downs
wd_keywords_file[0].close()
        
#%%Start of program
if __name__ == '__main__':        
    # TEST RO FILE:
    # ropathfile = 'C:/Users/changkk/OneDrive - Southern California Edison/CURRENT_CODE/RO_PARSER_NEW/TEST_ROS/D27_Seq4138989_TOLA_805963E_11032023.pdf' #NEW NO WIRE DOWN
    # ropathfile = 'C:/Users/changkk/OneDrive - Southern California Edison/CURRENT_CODE/RO_PARSER_NEW/TEST_ROS/D31_Seq1769_SPEEDWAY_4628970E_10292023.pdf' #NEW WIRE DOWN
    # ropathfile = 'C:/Users/changkk/OneDrive - Southern California Edison/CURRENT_CODE/RO_PARSER_NEW/TEST_ROS/D59_Seq6294_PYTHON_5138564_09262023.pdf'
    
    def parse_ro_form(pdf_file_to_parse):
    
    # GLOBAL FLAGS
    # FOR BOB'S ROS
        has_code_colon = 0
        has_plpole_dolly_diff = 0
        has_4wd_diff = 0
    

    # MAKE FILE NAME
        ro_comps = pdf_file_to_parse.split('/')
        ro_folder_path = '/'.join(ro_comps[:-1]) + '/'
        ro = ro_comps[-1]
    
    # OPEN PDF
        store_comps = []
    
    # CREATE PDF ELEMENT HIERARCHY BASED ON PHYSICAL POSITION IN DOC
    # DETERMINE NUMBER OF PAGES
        n_pages = len(list(PDFPage.get_pages(open(ro_folder_path + ro, 'rb'))))
        logger.info("Parsing %s (%s page(s))", ro, n_pages)
    # print(n_pages)
    
        for pg in range(n_pages):
            tmppg = [pg]
            for page_layout in extract_pages(ro_folder_path + ro, page_numbers=[pg]):
                for element in page_layout:
                    if isinstance(element, LTTextContainer):
                        elevars = vars(element)['_objs']
                        _ = [tmppg.append(str(ev)) for ev in elevars if "'\\n'>" not in str(ev)]
            store_comps.append(tmppg)
                
        comps_df = []
        for sc1 in store_comps:
            comps1 = sc1[1:]
            for sc in comps1:
                comps = sc.split(' ')[1:]
                pos = [float(c) for c in comps[0].split(',')]
                textdata = ' '.join(comps[1:]).replace('\\n','').replace('>','').replace("'",'')
                data = pos + [textdata]
                header = ['x0', 'y0', 'x1', 'y1', 'textdata']
                tmpdf = pd.DataFrame([data], columns=header)
                tmpdf['page'] = sc1[0]
                comps_df.append(tmpdf)
    
        comps_df = pd.concat(comps_df)
        comps_df = comps_df.reset_index(drop=True)
        comps_df = comps_df.sort_values(by=['page', 'y0', 'x0'], ascending=[True, False, True])
        comps_df = comps_df.reset_index(drop=True)
    # print(comps_df)
    
    # BECAUSE ONE FILE HAD A ADDITIONAL ATTACHMENTS, REMOVING ANYTHING
    # THAT MIGHT BE ATTACHED BEFORE THE WORDS "DISTRIBUTION REPAIR ORDER"
        dro_index = comps_df[comps_df['textdata'] == 'Distribution Repair Order'].index[0]
        comps_df = comps_df.iloc[dro_index:]
        comps_df = comps_df.reset_index(drop=True)
        comps_df = comps_df[~comps_df['textdata'].isin(['', ' '])]
    # print(comps_df)
    
    # SOME ROS SEEM TO RETAIN THE PDF FORMATTING STRINGS, REMOVE THEM HERE
    # AND REMOVE THE RO Ver
        comps_df = comps_df[~comps_df['textdata'].astype(str).str.startswith('matrix=[')]
        comps_df = comps_df[~comps_df['textdata'].astype(str).str.startswith('RO Ver')]
        comps_df = comps_df.reset_index(drop=True)
    # print(comps_df)
    
    # GROUP PDF ELEMENTS BY PHYSICAL POSITION TO IMPROVING PARSING
        comps_dict = dict()
        for (page_number, _y0), group in comps_df.groupby(['page', 'y0']):
            tmptext = ' '.join(group['textdata'].to_list())
            comps_dict.setdefault(page_number, []).append(tmptext)
    
        comps_list = []
        for k, v in comps_dict.items():
            comps_list.extend(v[::-1])
   
        material_idx = [idx for idx, x in enumerate(comps_list) if x == 'Material:'][0]
        information_list = comps_list[:material_idx]
        detail_list = comps_list[material_idx:]
    
        parsed_data = []
        parsed_data_headers = []
        for clidx, cl in enumerate(information_list):
            if "CAD ID:" in cl and 'Incident Date:' in cl:
                cadid_idx = cl.index('CAD ID:')
                incdate_idx = cl.index('Incident Date:')
                inctime_idx = cl.index('Time:')
            
                cadid = cl[cadid_idx+7:incdate_idx].strip()
                incdate = cl[incdate_idx+14:inctime_idx].strip().split('//')
                incdate = incdate[-1] + '-' + incdate[0] + '-' + incdate[1]
                inctime = cl[inctime_idx+5:].strip()
                if inctime.count(':') == 1:
                    inctime = inctime + ':00'
                incdt = incdate + ' ' + inctime
            # print(cadid, incdt)
            
                parsed_data.extend([cadid, incdt, incdate, inctime])
                parsed_data_headers.extend(['CAD_ID', 'INCIDENT_DATETIME', 'INCIDENT_DATE', 'INCIDENT_TIME'])
        
            elif "Structure No:" in cl and "RO Date:" in cl:
                structureno_idx = cl.index('Structure No:')
                rodate_idx = cl.index('RO Date:')
                rotime_idx = cl.index('Time:')
            
                structureno = cl[structureno_idx+13:rodate_idx].strip()
                rodate = cl[rodate_idx+8:rotime_idx].strip().split('//')
                rodate = rodate[-1] + '-' + rodate[0] + '-' + rodate[1]
                rotime = cl[rotime_idx+5:].strip()
                if rotime.count(':') == 1:
                    rotime = rotime + ':00'
                rodt = rodate + ' ' + rotime
            # print(structureno, rodt)
            
                parsed_data.extend([structureno, rodt])
                parsed_data_headers.extend(['STRUCTURE', 'RO_DATETIME'])
        
            elif "Reported by:" in cl and "Radio No:" in cl:
                reportedby_idx = cl.index('Reported by:')
                radiono_idx = cl.index('Radio No:')
                cellno_idx = cl.index('Cell Number:')
            
                reportedby = cl[reportedby_idx+12:radiono_idx].strip()
                radiono = cl[radiono_idx+9:cellno_idx].strip()
                cellno = cl[cellno_idx+12:].strip()
            # print('\n', reportedby, radiono, cellno, '\n')
        
                parsed_data.extend([reportedby, radiono, cellno])
                parsed_data_headers.extend(['REPORTED_BY', 'RADIO_NO', 'CELL_NO'])
        
            elif "Problem:" in cl and "Cause:" in cl:
                problem_idx = cl.index('Problem:')
                cause_idx = cl.index('Cause:')
            
                problem = cl[problem_idx+8:cause_idx].strip()
                cause = cl[cause_idx+6:].strip()
            # print('\n', problem, cause, '\n')
        
                parsed_data.extend([problem, cause])
                parsed_data_headers.extend(['PROBLEM', 'CAUSE'])
        
            elif "Circuit:" in cl and "Code" in cl:
                circuit_idx = cl.index('Circuit:')
                code_idx = cl.index('Code')
                sub_idx = cl.index('Sub:')
            
                circuit = cl[circuit_idx+8:code_idx].strip()
                try:
                    code_idx = cl.index('Code:')
                    code = cl[code_idx+5:sub_idx].strip()
                    has_code_colon = 1
                except:
                    code_idx = cl.index('Code')
                    code = cl[code_idx+4:sub_idx].strip()
                sub = cl[sub_idx+4:].strip()
            # print('\n', circuit, code, sub, '\n')
        
                parsed_data.extend([circuit, code, sub])
                parsed_data_headers.extend(['CIRCUIT', 'CODE', 'SUB'])
        
            elif "Location:" in cl:
                location = cl.replace('Location:','').strip()
                location_city = location[location.lower().rindex(',')+1:].strip() if ',' in location.lower() else ''
            # print('\n', location, '\n')
        
                parsed_data.extend([location, location_city])
                parsed_data_headers.extend(['LOCATION', 'LOCATION_CITY'])
        
            elif "Access:" in cl:
                access = cl.replace('Access:', '').strip()
            # print('\n', access, '\n')
        
                parsed_data.extend([access])
                parsed_data_headers.extend(['ACCESS'])
        
            elif "Seq No:" in cl and "Foreman:" in cl:
                seqno_idx = cl.index('Seq No:')
                foreman_idx = cl.index('Foreman:')
                radiono_idx2 = cl.index('Radio No:')
            
                seqno = cl[seqno_idx+7:foreman_idx].strip()
                foreman = cl[foreman_idx+8:radiono_idx2].strip()
            # print('\n', seqno, foreman, '\n')
        
                parsed_data.extend([seqno, foreman])
                parsed_data_headers.extend(['SEQ_NO', 'FOREMAN'])
        
            elif "Duration" in cl and "Crew Size:" in cl:
                duration_idx = cl.index('Duration:')
                crewsize_idx = cl.index('Crew Size:')
                crewsneeded_idx = cl.index('Crews Needed:')
                crewtype_idx = cl.index('Crew Type:')
            
                duration = cl[duration_idx+9:crewsize_idx].strip()
                crewsize = cl[crewsize_idx+10:crewsneeded_idx].strip()
                crewsneeded = cl[crewsneeded_idx+13:crewtype_idx].strip()
                crewtype = cl[crewtype_idx+10:].strip()
            # print('\n', duration, crewsize, crewsneeded, crewtype, '\n')
            
                parsed_data.extend([duration, crewsize, crewsneeded, crewtype])
                parsed_data_headers.extend(['DURATION', 'CREW_SIZE', 'CREWS_NEEDED', 'CREW_TYPE'])
            
            elif "Enviromental Release?" in cl and "Enviromental Other?" in cl:
                envrel_idx = cl.index('Enviromental Release?')
                envoth_idx = cl.index('Enviromental Other')
                distnot_idx = cl.index('District Notified?')
            
                envrel = cl[envrel_idx+21:envoth_idx].strip()
                envoth = cl[envoth_idx+19:distnot_idx].strip()
                distnot = cl[distnot_idx+18:].strip()
            # print('\n', envrel, envoth, distnot, '\n')
        
                parsed_data.extend([envrel, envoth, distnot])
                parsed_data_headers.extend(['ENVIRONMENTAL_RELEASE', 'ENVIRONMENTAL_OTHER', 'DISTRICT_NOTIFIED'])
        
            elif "Tree Crew Needed?" in cl:
                treecrew_idx = cl.index('Tree Crew Needed?')
                treecrewgnd_idx = cl.index('Grounding Needed for Tree Crew?')
            
                treecrewneeded = cl[treecrew_idx+17:treecrewgnd_idx].strip()
                treecrewgrounding = cl[treecrewgnd_idx+31:].strip()
            # print('\n', treecrewneeded, treecrewgrounding, '\n')
        
                parsed_data.extend([treecrewneeded, treecrewgrounding])
                parsed_data_headers.extend(['TREE_CREW_NEEDED', 'GROUNDING_NEEDED_FOR_TREE_CREW'])
        
            elif 'Load Off (Individual or Downstream):' in cl:
                loadoff = information_list[clidx+1].strip()
            # print('\n', loadoff, '\n')

                parsed_data.extend([loadoff])
                parsed_data_headers.extend(['LOAD_OFF'])

            elif 'Other Departments:' in cl:
                depts = information_list[clidx+1].strip()
                telecom_idx = depts.index('TeleComm')
                trans_idx = depts.index('Transmission')
                util_idx = depts.index('Other Utility')
                oth_idx = depts.index('Other')
            
                telecom = depts[:telecom_idx].strip()
                transmission = depts[telecom_idx+8:trans_idx].strip()
                utility = depts[trans_idx+12:util_idx].strip()
                other = depts[util_idx+13:oth_idx].strip()
            # print('\n', telecom, transmission, utility, other, '\n')
        
                parsed_data.extend([telecom, transmission, utility, other])
                parsed_data_headers.extend(['OTHDEPT_TELECOM', 
                                        'OTHDEPT_TRANSMISSION', 
                                        'OTHDEPT_OTHER_UTILIITY', 
                                        'OTHDEPT_OTHER'])
        
            elif 'Equipment Needed:' in cl:
                #GET INDEX OF 'Equipment Needed:' TEXT FIRST
                eqneededidx = information_list.index('Equipment Needed:')
                
                eq1 = information_list[clidx+1]
                eq2 = information_list[clidx+2]
                eq3 = information_list[clidx+3]
                eq4 = information_list[clidx+4]

                dbltub_idx = eq1.index('Double/Tub Bucket')
                
                # TO ACCOUNT FOR 'Traffic Control' INSTEAD OF 'Arrowboard'
                # IN NEWER ROs
                try:
                    arrow_idx = eq1.index('Arrowboard')
                except:
                    arrow_idx = eq1.index('Traffic Control')
                # arrow_idx = eq1.index('Arrowboard')
                flat_idx = eq1.index('Flatbed')
                easy_idx = eq1.index('EasyHauler')
                dbltub = eq1[:dbltub_idx].strip()
                arrow = eq1[dbltub_idx+17:arrow_idx].strip()
                flat = eq1[arrow_idx+10:flat_idx].strip()
                easy = eq1[flat_idx+7:easy_idx].strip()
            
                y2k_idx = eq2.index('Y2K')
                pedcont_idx = eq2.index('Ped Control')
                snow_idx = eq2.index('Snow Cat')
                fork_idx = eq2.index('Fork Lift')
                y2k = eq2[:y2k_idx].strip()
                pedcont = eq2[y2k_idx+3:pedcont_idx].strip()
                snow = eq2[pedcont_idx+11:snow_idx].strip()
                fork = eq2[snow_idx+8:fork_idx].strip()            
        
                auger_idx = eq3.index('Auger Truck')
                pldolly_idx = eq3.index('P/L Dolly')
                air_idx = eq3.index('Air Ops')
                pwrdolly_idx = eq3.index('Power Dolly')
                auger = eq3[:auger_idx].strip()
                pldolly = eq3[auger_idx+11:pldolly_idx].strip()
                air = eq3[pldolly_idx+9:air_idx].strip()
                pwrdolly = eq3[air_idx+7:pwrdolly_idx].strip()
        
                ug_idx = eq4.index('UG Van')
                try:
                    plpoledolly_idx = eq4.index('P/L Pole Dolly')
                except:
                    plpoledolly_idx = eq4.index('P/LPole  Dolly')
                    has_plpole_dolly_diff = 1
                try:
                    fourwd_idx = eq4.index('Four Wheel Drive')
                except:
                    fourwd_idx = eq4.index('FWD')
                    has_4wd_diff = 1
                crane_idx = eq4.index('Crane')
                ug = eq4[:ug_idx].strip()
                plpoledolly = eq4[ug_idx+6:plpoledolly_idx].strip()
                fourwd = eq4[plpoledolly_idx+14:fourwd_idx].strip()
                crane = eq4[fourwd_idx+16:crane_idx].strip()
        
            # print('\n', dbltub, arrow, flat, easy,
            #             y2k, pedcont, snow, fork,
            #             auger, pldolly, air, pwrdolly,
            #             ug, plpoledolly, fourwd, crane, '\n')
        
                parsed_data.extend([dbltub, arrow, flat, easy,
                                y2k, pedcont, snow, fork,
                                auger, pldolly, air, pwrdolly,
                                ug, plpoledolly, fourwd, crane])
            
                parsed_data_headers.extend(['EQ_DOUBLE_TUB/BUCKET', 'EQ_ARROWBOARD', 'EQ_FLATBED', 'EQ_EASYHAULER',
                                        'EQ_Y2K', 'EQ_PED_CONTROL', 'EQ_SNOW_CAT', 'EQ_FORK_LIFT',
                                        'EQ_AUGER_TRUCK', 'EQ_P/L_DOLLY', 'EQ_AIR_OPS', 'EQ_POWER_DOLLY',
                                        'EQ_UG_VAN', 'EQ_P/L_POLE_DOLLY', 'EQ_FOUR_WHEEL_DRIVE', 'EQ_CRANE'])
        
        # print(cl)
                
        for dlidx, dl in enumerate(detail_list):
            if "# of Attachments-" in dl or "# of Attachments -" in dl:
                if has_code_colon and has_plpole_dolly_diff and has_4wd_diff:
                    endidx = dlidx-1
                else:
                    endidx = dlidx
            
                materials = '\n'.join(detail_list[1:endidx+1])
            # print(materials)
        
                parsed_data.extend([materials])
                parsed_data_headers.extend(['MATERIALS'])
            
            
                n_attachments = detail_list[endidx].replace('# of Attachments-','').replace('# of Attachments -','').strip()                
                       
                parsed_data.extend([n_attachments])
                parsed_data_headers.extend(['N_ATTACHMENTS'])

            elif "Remarks:(" in dl:
            # CHECK FOR "See Remarks Attachment
                check_remarks_attachment = [dlidx2 for dlidx2, dl2 in enumerate(detail_list) if 'See Remarks Attachment' in dl2]
                if len(check_remarks_attachment) > 0:                    
                    remattidx = [dlidx2 for dlidx2, dl2 in enumerate(detail_list) if 'See Remarks Attachment' in dl2][0]
                    pgidx = [dlidx2 for dlidx2, dl2 in enumerate(detail_list) if 'Pg 2 of ' in dl2][0]
                    remarks1 = detail_list[dlidx+1:remattidx]
                    remarks2 = detail_list[remattidx+3:pgidx]
                    remarks = remarks1 + remarks2
                    remarks = '\n'.join(remarks)
                # print(remarks)
            
                    parsed_data.extend([remarks])
                    parsed_data_headers.extend(['REMARKS'])
            
                else:
                    try:
                        pgidx = [dlidx2 for dlidx2, dl2 in enumerate(detail_list) if 'Pg 1 of ' in dl2][0]
                    except:
                        pgidx = len(detail_list)
                        
                    # CHANGED DUE TO ODD CHANGES IN THE WAY SOME ROS HAD BEEN STRUCTURED
                    
                    # remarksandwd = detail_list[dlidx+1:pgidx]
                    remarksandwd = detail_list[dlidx+1:]
                # remarks = '\n'.join(detail_list[dlidx+1:pgidx])
                # print(remarks)
                
                # ADD FIELDS FOR NEW ROS WITH WIRE DOWN INFORMATION
                # FOR ROS WITHOUT WIRE DOWN INFORMATION, JUST NAN THE FIELD
                # VALUES
                    try:
                        wdinfoheaderidx = remarksandwd.index('Wire Down Information')
                    # SPLIT INTO REGULAR REMARKS AND WIRE DOWN INFO
                        remarks = '\n'.join(remarksandwd[:wdinfoheaderidx])
                        parsed_data.extend([remarks])
                        parsed_data_headers.extend(['REMARKS'])
                    
                        wdfields = remarksandwd[wdinfoheaderidx+1:]
                        wdfieldsdf = pd.DataFrame(wdfields,
                                              columns=['WD_FIELDS'])
                        wdfieldsdf2 = wdfieldsdf['WD_FIELDS'].astype(str).str.split(':', expand=True)
                        wdfieldsdf2.columns = ['WD_FIELDS', 'WD_VALUES']
                        wdfieldsdf2['WD_FIELDS'] = 'WD_' + wdfieldsdf2['WD_FIELDS'].astype(str)
                        wdfieldsdf2['WD_VALUES'] = wdfieldsdf2['WD_VALUES'].astype(str).str.strip()
                    
                        for wdidx, wdrow in wdfieldsdf2.iterrows():
                           parsed_data.extend([wdrow['WD_VALUES']])
                           parsed_data_headers.extend([wdrow['WD_FIELDS']])
                
                    except:
                        remarks = '\n'.join(remarksandwd)
                        parsed_data.extend([remarks])
                        parsed_data_headers.extend(['REMARKS'])
                    
                    # NaN OUT THE WD FIELDS IF RO DOESN'T HAVE WD SECTION.
                        wd_fields = ['WD_Reported by', 
                         'WD_Radio No', 
                         'WD_Cell Number', 
                         'WD_Structure nearest to the wire-down incident( Structure number)', 
                         'WD_Was the wire down energized unpon arrival ?', 
                         'WD_Was the wire down deenergized Manually prior to arrival ? (MADEC, TSO Opened)', 
                         'WD_Line or Load Side Wire Down?', 
                         'WD_Was the wire down covered conductor?', 
                         'WD_What was the physical condition of the line?', 
                         'WD_Where did it fall?', 
                         'WD_What contributed to the incident?', 
                         'WD_Was the line section fused?', 
                         'WD_Fuse operation?', 
                         'WD_Was there a hotline down from backfeed?']
                    
                        for wdf in wd_fields:
                            parsed_data.extend([np.nan])
                            parsed_data_headers.extend([wdf])
                    
                
    # CHECK IF MATERIALS AND REMARKS IN COLUMNS - WON'T BE IF THERE
    # IS A PARSING ERROR
        if 'MATERIALS' not in parsed_data_headers:
            parsed_data.extend(['PARSE_ERROR'])
            parsed_data_headers.extend(['MATERIALS'])
        
        if 'N_ATTACHMENTS' not in parsed_data_headers:
            parsed_data.extend(['PARSE_ERROR'])
            parsed_data_headers.extend(['N_ATTACHMENTS'])
    
        if 'REMARKS' not in parsed_data_headers:
            parsed_data.extend(['PARSE_ERROR'])
            parsed_data_headers.extend(['REMARKS'])            
        
    # MAKE FINAL RO TEXT DATAFRAME:
        rodf = pd.DataFrame([parsed_data], columns=parsed_data_headers)
        rodf['FILENAME'] = ro
        
# Search for keywords in PDF text
        reader_wd = PdfReader(pdf_file_to_parse)
        page_wd = reader_wd.pages[0]
        text_wd = page_wd.extract_text().strip()
        header_index_wd = text_wd.lower().index('distribution repair order')
        text_header_wd = text_wd[header_index_wd+len('distribution repair order'):]
        wd_match = False
        wd_matched_keyword = ''
        
# Search for WD notes
        event_notes = ''
        if 'wd call:' in text_header_wd.lower():
            event_notes = text_parser(text_header_wd, 'wd call:', '\nmaterial:').strip()

        for wd_keyword in wd_keywords:
            if wd_keyword in text_wd.lower():
                wd_match = True
                wd_matched_keyword = wd_keyword
                break
        
        rodf['wd_match'] = wd_match
        rodf['wd_matched_keyword'] = wd_matched_keyword
        rodf['WD_Event_Notes'] = event_notes
        
        
        rodf = rodf.replace(' ', np.nan).replace('', np.nan)
        ro_append.append(rodf)
        logger.debug("Parsed %s; wire-down keyword match: %s", ro, wd_matched_keyword or "none")
        # print(rodf)
    

    site_url = os.environ.get(
        'RO_SHAREPOINT_SITE_URL',
        'https://edisonintl.sharepoint.com/teams/td3/AES/AADS/FA/Failure-Event-Analysis',
    )
    library_name = os.environ.get('RO_SHAREPOINT_LIBRARY', 'Documents')
    dir_path = os.environ.get(
        'RO_SHAREPOINT_FOLDER', 'Repair Order Event Tracking/Keyword Scrape',
    )
    shared_drive_path = '//sce/workgroup/TDBU7/RESO-CRT/PERFORMANCE MANAGEMENT AND ANALYSIS/RO_PDFs'
    
    temporary_directory = TemporaryDirectory(prefix='ro_pdf_scraper_')
    try:
        ro_append = []

        auth_path = str(PROJECT_ROOT / 'config' / 'sharepoint_auth.json')
        graph = GraphClient(app_config_path=auth_path, credential_file=auth_path)
        drive_endpoint, folder_endpoint, new_folder_id, old_folder_id = resolve_sharepoint_folder(
            graph, site_url, library_name, dir_path,
        )
        pdf_items = [
            item for item in graph.get_all(f"{drive_endpoint}/items/{new_folder_id}/children")
            if 'file' in item and item['name'].lower().endswith('.pdf')
        ]
        logger.info("Starting repair-order PDF run; SharePoint source: %s / %s / %s/New",
                    site_url, library_name, dir_path)
        logger.info("Found %s PDF file(s) to process", len(pdf_items))

        if pdf_items:
            for item in pdf_items:
                pdf_name = item['name']
                try:
                    pdf_file_to_parse = (Path(temporary_directory.name) / pdf_name).as_posix()
                    item_endpoint = f"{drive_endpoint}/items/{item['id']}"
                    download_sharepoint_file(graph, item_endpoint, pdf_file_to_parse)
                    parse_ro_form(pdf_file_to_parse)
                
        # Create a copy in Shared Drive and also Move pdfs to another folder, clearing out original folder
                    shutil.copy2(pdf_file_to_parse, shared_drive_path + '/' + pdf_name)
                    graph.patch(item_endpoint, {'parentReference': {'id': old_folder_id}})
                           
                    logger.info("Processed and archived %s", pdf_name)
                    
                except Exception as e:
                #If an error occurs, print the error and skip this file
                    logger.exception("Error processing %s: %s", pdf_name, e)
                    continue
            
            ro_df_final = pd.concat(ro_append) if ro_append else pd.DataFrame()
            
            
        else:
            ro_df_final = []
            logger.info("No repair-order PDFs found")
            
        if len(ro_df_final) > 0:
            
            ro_df_final['INCIDENT_DATE'] = pd.to_datetime(ro_df_final['INCIDENT_DATE']).dt.strftime('%m/%d/%Y')
            ro_df_final['INCIDENT_TIME'] = pd.to_datetime(ro_df_final['INCIDENT_TIME']).dt.strftime('%H:%M')
        # Open xlsx file
            workbook_item = graph.get(f"{folder_endpoint}:/Repair%20Order%20Data.xlsx")
            workbook_endpoint = f"{drive_endpoint}/items/{workbook_item['id']}"
            output_file = Path(temporary_directory.name) / 'Repair Order Data.xlsx'
            download_sharepoint_file(graph, workbook_endpoint, output_file)
            output_wb = load_workbook(output_file)
            output_ws = output_wb['Sheet1']
            output_wb.active = output_ws

        #create data to write to excel on SharePoint for PowerAutomate to pick up (Need code updated to pull row['Statement_of_Facts'] 
        #   and row['Facility Identification'] and then change excel on SharePoint to have columns also)
            for index, row in ro_df_final.iterrows():
                write_data = [row['FILENAME'], row['REMARKS'], row['INCIDENT_DATE'], row['INCIDENT_TIME'], row['PROBLEM'], 
                              row['CAUSE'], row['CAD_ID'], row['STRUCTURE'], row['CIRCUIT'], row['wd_match'], row['wd_matched_keyword'],
                              row['WD_Structure nearest to the wire-down incident( Structure number)'], row['WD_Was the wire down energized unpon arrival ?'],
                              row['WD_Line or Load Side Wire Down?'], row['WD_Was the wire down covered conductor?'], row['SEQ_NO'], row['RADIO_NO'],
                              row['CREW_TYPE'], row['LOCATION'], row['LOCATION_CITY'], row['MATERIALS'], row['WD_Event_Notes']]

                output_ws.append(write_data)

            output_ws.tables['Data'].ref = "A1:" + output_ws.cell(output_ws.max_row-1, output_ws.max_column).column_letter + \
                                                str(output_ws.max_row)

            output_wb.save(output_file)
            output_wb.close()
            upload_sharepoint_workbook(graph, workbook_endpoint, output_file)
            logger.info("Saved %s parsed row(s) to SharePoint workbook", len(ro_df_final))
        logger.info('Completed running RO_PDF_PARSER_SC.py')
    except Exception as e:
        logger.exception('RO_PDF_PARSER_SC.py failed to run: %s', e)
    finally:
        temporary_directory.cleanup()
        
    